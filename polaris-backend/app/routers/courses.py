"""课程表路由：学期管理、课程 CRUD、批量导入、冲突检测、课表查询。"""
from __future__ import annotations

from fastapi import APIRouter, Depends, File, Query, UploadFile
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.deps import get_current_user
from app.core.exceptions import AppError
from app.models.user import User
from app.schemas.course import (CourseImportIn, CourseIn, CourseUpdate, SemesterIn, SemesterUpdate)
from app.services import course_service

router = APIRouter(prefix="/courses", tags=["② 课程表"])


# ─────────── 学期管理 ───────────
@router.get("/semesters", summary="学期列表")
def list_semesters(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return {"items": [s.to_dict() for s in course_service.list_semesters(db, user.id)]}


@router.post("/semesters", summary="新建学期")
def create_semester(body: SemesterIn, db: Session = Depends(get_db),
                    user: User = Depends(get_current_user)) -> dict:
    from app.models.course import Semester

    if body.is_current:
        db.query(Semester).filter(Semester.user_id == user.id).update({"is_current": False})
    semester = Semester(user_id=user.id, name=body.name, start_date=body.start_date,
                        end_date=body.end_date, total_weeks=body.total_weeks,
                        is_current=body.is_current)
    db.add(semester)
    db.commit()
    db.refresh(semester)
    return semester.to_dict()


@router.put("/semesters/{semester_id}", summary="更新学期")
def update_semester(semester_id: int, body: SemesterUpdate, db: Session = Depends(get_db),
                    user: User = Depends(get_current_user)) -> dict:
    from app.models.course import Semester

    semester = db.get(Semester, semester_id)
    if not semester or semester.user_id != user.id:
        raise AppError("学期不存在", code="not_found", status_code=404)
    if body.is_current:
        db.query(Semester).filter(Semester.user_id == user.id).update({"is_current": False})
    for key, value in body.model_dump(exclude_unset=True).items():
        if value is not None:
            setattr(semester, key, value)
    db.commit()
    db.refresh(semester)
    return semester.to_dict()


@router.post("/semesters/{semester_id}/current", summary="设为当前学期")
def set_current(semester_id: int, db: Session = Depends(get_db),
                user: User = Depends(get_current_user)) -> dict:
    return course_service.set_current_semester(db, user.id, semester_id).to_dict()


@router.delete("/semesters/{semester_id}", summary="删除学期（课程归属置空）")
def delete_semester(semester_id: int, db: Session = Depends(get_db),
                    user: User = Depends(get_current_user)) -> dict:
    from app.models.course import Semester

    semester = db.get(Semester, semester_id)
    if not semester or semester.user_id != user.id:
        raise AppError("学期不存在", code="not_found", status_code=404)
    db.delete(semester)
    db.commit()
    return {"deleted": True}


# ─────────── 课表视图（放在 /{course_id} 之前） ───────────
@router.get("/week", summary="周课表（含当前教学周）")
def week_timetable(week: int | None = Query(default=None), semester_id: int | None = Query(default=None),
                   db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return course_service.week_timetable(db, user.id, week=week, semester_id=semester_id)


@router.get("/today", summary="今日课程")
def today(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    classes = course_service.today_classes(db, user.id)
    return {"date": classes[0]["date"] if classes else None, "total": len(classes), "classes": classes}


@router.post("/conflicts", summary="冲突预检（不落库，用于前端实时提示）")
def check_conflicts(body: CourseIn, db: Session = Depends(get_db),
                    user: User = Depends(get_current_user)) -> dict:
    conflicts = course_service.find_conflicts(db, user.id, body.schedules,
                                              semester_id=body.semester_id)
    return {"has_conflict": bool(conflicts), "count": len(conflicts), "conflicts": conflicts}


# ─────────── 批量导入 ───────────
@router.post("/import", summary="批量导入课程（JSON）")
def import_json(body: CourseImportIn, db: Session = Depends(get_db),
                user: User = Depends(get_current_user)) -> dict:
    return course_service.import_courses(db, user.id, body)


@router.post("/import-excel", summary="批量导入课程（Excel：课程名/老师/教室/星期/开始/结束/周次/学分）")
async def import_excel(file: UploadFile = File(...), semester_id: int | None = Query(default=None),
                       force: bool = Query(default=False),
                       db: Session = Depends(get_db),
                       user: User = Depends(get_current_user)) -> dict:
    data = await file.read()
    if not data:
        raise AppError("Excel 文件为空")
    raw_courses = course_service.parse_course_excel(data)
    if not raw_courses:
        raise AppError("未从 Excel 中解析到有效课程行")
    payload = CourseImportIn(semester_id=semester_id, replace=force, courses=raw_courses)
    return course_service.import_courses(db, user.id, payload)


@router.get("/import-template", summary="下载课表 Excel 导入模板（平铺表格式）")
def import_template(user: User = Depends(get_current_user)):
    import io

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "课表导入"
    headers = ["课程名", "老师", "教室", "星期", "开始时间", "结束时间", "周次", "学分"]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1F6FEB")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
    ws.append(["高等数学", "王老师", "信东201", "周一", "08:00", "09:40", "1-16", 4])
    ws.append(["大学英语", "李老师", "教三305", "周二", "10:00", "11:40", "1-16双", 2])
    ws.append(["操作系统", "张老师", "信东101", "周三", "14:00", "15:40", "1-12", 3])
    for col, w in zip("ABCDEFGH", [16, 10, 12, 8, 11, 11, 10, 6]):
        ws.column_dimensions[col].width = w
    ws["A10"] = "填写说明：① 星期可填 周一/星期二/1~7；② 时间 HH:MM；③ 周次如 1-16、1-16单、1-16双、1,3,5；"
    ws["A11"] = "④ 同一门课多天安排请重复行（课程名相同自动合并）；⑤ 也支持直接上传学校官方网格课表 Excel。"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=courses_import_template.xlsx"})


@router.delete("", summary="清空课程（可按学期范围；用于撤销错误导入）")
def clear_courses(semester_id: int | None = Query(default=None),
                  db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    from app.models.course import Course, CourseSchedule

    # 只清空个人课表；教室课表请到空教室页清空
    q = db.query(Course).filter(Course.user_id == user.id, Course.source == "personal")
    if semester_id:
        q = q.filter(Course.semester_id == semester_id)
    ids = [c.id for c in q.all()]
    db.query(CourseSchedule).filter(CourseSchedule.course_id.in_(ids)).delete(synchronize_session=False)
    n = q.delete(synchronize_session=False)
    db.commit()
    return {"deleted": n}


# ─────────── 课程 CRUD ───────────
@router.get("", summary="课程列表（可按学期/关键词筛选）")
def list_courses(semester_id: int | None = Query(default=None), keyword: str | None = Query(default=None),
                 db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    courses = course_service.list_courses(db, user.id, semester_id=semester_id, keyword=keyword)
    return {"total": len(courses), "items": [c.to_dict() for c in courses]}


@router.post("", summary="新建课程（force=true 可忽略冲突强制保存）")
def create_course(body: CourseIn, force: bool = Query(default=False),
                  db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    course = course_service.create_course(db, user.id, body, force=force)
    return course.to_dict()


@router.get("/{course_id}", summary="课程详情")
def get_course(course_id: int, db: Session = Depends(get_db),
               user: User = Depends(get_current_user)) -> dict:
    return course_service.get_course(db, user.id, course_id).to_dict()


@router.put("/{course_id}", summary="更新课程（schedules 传入则整体替换）")
def update_course(course_id: int, body: CourseUpdate, force: bool = Query(default=False),
                  db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return course_service.update_course(db, user.id, course_id, body, force=force).to_dict()


@router.delete("/{course_id}", summary="删除课程")
def delete_course(course_id: int, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)) -> dict:
    course_service.delete_course(db, user.id, course_id)
    return {"deleted": True}
