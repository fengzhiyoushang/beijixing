"""空教室路由：教室管理、状态上报、空闲概率、预测推荐、VL 识图、使用记录导入。"""
from __future__ import annotations

import io
from datetime import date

from fastapi import APIRouter, Depends, File, Form, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.deps import get_current_user
from app.core.exceptions import AppError
from app.models.user import User
from app.schemas.classroom import ClassroomIn, ClassroomUpdate, StatusReportIn, VisionRecognizeIn
from app.services import classroom_service, storage

router = APIRouter(prefix="/classroom", tags=["④⑤ 空教室"])


# ─────────── 分析类（置于 /classrooms/{id} 之前） ───────────
@router.get("/buildings", summary="教学楼与教室清单")
def buildings(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return {"items": classroom_service.list_buildings(db)}


@router.get("/overview", summary="各教学楼空闲概况")
def overview(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return {"items": classroom_service.building_overview(db)}


@router.get("/free-rate", summary="某教室的空闲率矩阵（weekday × hour）")
def free_rate(building: str = Query(...), room_no: str = Query(...),
              days_back: int = Query(default=120, ge=7, le=730),
              db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return classroom_service.free_rate(db, building=building, room_no=room_no,
                                       days_back=days_back, user_id=user.id)


@router.get("/building-usage", summary="某教学楼每间教室的使用状态图（绿=空闲/红=占用/灰=无数据）")
def building_usage(building: str = Query(...), days_back: int = Query(default=120, ge=7, le=730),
                   db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return classroom_service.building_usage(db, building, days_back=days_back, user_id=user.id)


@router.get("/campus-usage", summary="全校教学楼使用状态（每楼一张 楼层×教室序号 图的数据源）")
def campus_usage(days_back: int = Query(default=120, ge=7, le=730),
                 db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return classroom_service.campus_usage(db, days_back=days_back, user_id=user.id)


@router.get("/predict", summary="空闲教室预测推荐（时间衰减加权 + 课程占用惩罚）")
def predict(weekday: int | None = Query(default=None, ge=1, le=7),
            hour: int | None = Query(default=None, ge=0, le=23),
            days_back: int = Query(default=90, ge=7, le=730),
            limit: int = Query(default=10, ge=1, le=50),
            db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return classroom_service.predict(db, user.id, weekday=weekday, hour=hour,
                                     days_back=days_back, limit=limit)


@router.get("/today", summary="今日采集概况")
def today(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return classroom_service.today_stats(db)


# ─────────── 状态上报 ───────────
@router.post("/status", summary="上报教室状态（Web/小程序手动）")
def report_status(body: StatusReportIn, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)) -> dict:
    from app.services import dashboard_service

    log = classroom_service.report_status(db, user.id, body)
    dashboard_service.invalidate(user.id)
    return log.to_dict()


@router.post("/status/photo", summary="上报教室状态（multipart 图片，含照片存证）")
async def report_with_photo(
    building: str = Form(...),
    room_no: str = Form(...),
    status: str = Form(...),
    note: str | None = Form(default=None),
    occupied_seats: int = Form(default=0),
    photo: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict:
    data = await photo.read()
    saved = storage.save_upload("classroom", photo.filename or "room.jpg", data)
    payload = StatusReportIn(building=building, room_no=room_no, status=status, note=note,
                             occupied_seats=occupied_seats, source="miniapp",
                             photo_url=saved["url"], confidence=1.0)
    return classroom_service.report_status(db, user.id, payload).to_dict()


@router.get("/status/recent", summary="最近的状态上报记录")
def recent(building: str | None = Query(default=None), room_no: str | None = Query(default=None),
           limit: int = Query(default=20, ge=1, le=200),
           offset: int = Query(default=0, ge=0, le=100000),
           db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    logs = classroom_service.recent_logs(db, building=building, room_no=room_no,
                                         limit=limit, offset=offset)
    total_all = classroom_service.count_recent_logs(db, building=building, room_no=room_no)
    return {"total": total_all, "offset": offset, "items": [x.to_dict() for x in logs]}


@router.delete("/status/{log_id}", summary="删除一条教室状态采集记录")
def delete_status(log_id: int, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)) -> dict:
    return classroom_service.delete_status_log(db, log_id, user.id)


# ─────────── 多模态识图 ───────────
@router.post("/recognize", summary="AI 识图判定教室状态（base64 → deepseek-vl）")
async def recognize(body: VisionRecognizeIn, db: Session = Depends(get_db),
                    user: User = Depends(get_current_user)) -> dict:
    return await classroom_service.recognize_image_async(db, user.id, body)


@router.post("/recognize/upload", summary="AI 识图判定教室状态（直接上传图片文件）")
async def recognize_upload(
    photo: UploadFile = File(...),
    building: str | None = Form(default=None),
    room_no: str | None = Form(default=None),
    save_log: bool = Form(default=True),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict:
    import base64

    data = await photo.read()
    if not data:
        raise AppError("图片内容为空")
    payload = VisionRecognizeIn(
        building=building, room_no=room_no, save_log=save_log,
        image_base64=base64.b64encode(data).decode(),
    )
    return await classroom_service.recognize_image_async(db, user.id, payload)


# ─────────── 教室课表 Excel 导入（全校占用数据，与个人课表分离）───────────
def _clear_classroom_courses(db: Session, user_id: int, codes: set[str] | None = None) -> int:
    """清空教室课表；codes 给定时只删除排课地点命中这些教室编码的课程（按教室粒度替换）。"""
    from app.models.course import Course, CourseSchedule
    from app.services.classroom_excel_service import codes_in_text

    q = db.query(Course).filter(Course.user_id == user_id, Course.source == "classroom")
    courses = q.all()
    if codes:
        courses = [c for c in courses
                   if codes & {x for s in c.schedules for x in codes_in_text(s.location or c.location or "")}]
    if not courses:
        return 0
    ids = [c.id for c in courses]
    db.query(CourseSchedule).filter(CourseSchedule.course_id.in_(ids)).delete(synchronize_session=False)
    n = db.query(Course).filter(Course.id.in_(ids)).delete(synchronize_session=False)
    db.commit()
    return n


@router.post("/schedule/import-excel", summary="导入教室课表 Excel（支持多选文件；按教室粒度替换旧数据）")
async def schedule_import_excel(files: list[UploadFile] = File(...),
                                semester_id: int | None = Query(default=None),
                                db: Session = Depends(get_db),
                                user: User = Depends(get_current_user)) -> dict:
    """批量导入多个教室课表文件：每个文件涉及的教室，其旧数据被替换，其他教室数据保留。"""
    from app.schemas.course import CourseImportIn
    from app.services import classroom_excel_service, course_service, dashboard_service

    raw_courses = []
    touched_codes: set[str] = set()
    file_errors = []
    for f in files:
        name = f.filename or "未命名.xlsx"
        if not name.lower().endswith((".xlsx", ".xls")):
            file_errors.append(f"「{name}」不是 Excel 文件，已跳过")
            continue
        data = await f.read()
        if len(data) > 10 * 1024 * 1024:
            file_errors.append(f"「{name}」超过 10MB，已跳过")
            continue
        try:
            parsed = course_service.parse_course_excel(data)
        except Exception as exc:
            file_errors.append(f"「{name}」解析失败：{exc}")
            continue
        if not parsed:
            file_errors.append(f"「{name}」未解析到有效课程行")
            continue
        raw_courses.extend(parsed)
        for c in parsed:
            touched_codes.update(classroom_excel_service.codes_in_text(c.get("location") or ""))
            for s in c.get("schedules") or []:
                touched_codes.update(classroom_excel_service.codes_in_text(s.get("location") or ""))
    if not raw_courses:
        raise AppError("；".join(file_errors) or "未解析到任何有效课程")
    cleared = _clear_classroom_courses(db, user.id, touched_codes or None)
    payload = CourseImportIn(semester_id=semester_id, replace=False, courses=raw_courses)
    result = course_service.import_courses(db, user.id, payload, source="classroom")
    result["cleared"] = cleared
    result["rooms"] = sorted(touched_codes)
    result["file_errors"] = file_errors
    dashboard_service.invalidate(user.id)
    return result


@router.get("/schedule", summary="已导入的教室课表清单")
def schedule_list(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    from app.models.course import Course

    items = db.query(Course).filter(Course.user_id == user.id,
                                    Course.source == "classroom").order_by(Course.name).all()
    rooms = {r for c in items for s in c.schedules
             for r in [(s.location or c.location or "").strip()] if r}
    return {
        "total": len(items),
        "total_slots": sum(len(c.schedules) for c in items),
        "room_locations": len(rooms),
        "items": [{"id": c.id, "name": c.name, "teacher": c.teacher,
                   "slots": len(c.schedules), "location": c.location,
                   "created_at": c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else None}
                  for c in items],
    }


@router.post("/schedule/parse", summary="教室课表 Excel 解析检查（批量，可导出 xlsx/json）")
async def schedule_parse(files: list[UploadFile] = File(...),
                         export: str | None = Query(default=None, pattern="^(xlsx|json)$"),
                         user: User = Depends(get_current_user)):
    """解析多个教室课表 Excel：提取每格教室编码、拆解楼号/楼层/序号、校验文件名-标题-格子对应关系。

    export=xlsx|json 时直接返回文件下载；否则返回结构化 JSON 预览。"""
    from app.services import classroom_excel_service

    payload: list[tuple[str, bytes]] = []
    for f in files:
        name = f.filename or "未命名.xlsx"
        if not name.lower().endswith((".xlsx", ".xls")):
            raise AppError(f"「{name}」不是 Excel 文件，请上传 .xlsx")
        data = await f.read()
        if len(data) > 10 * 1024 * 1024:
            raise AppError(f"「{name}」超过 10MB 限制")
        payload.append((name, data))
    result = classroom_excel_service.parse_schedule_files(payload)
    if export == "json":
        return StreamingResponse(
            io.BytesIO(classroom_excel_service.build_export_json(result)),
            media_type="application/json",
            headers={"Content-Disposition": "attachment; filename=classroom_schedule_parse.json"})
    if export == "xlsx":
        return StreamingResponse(
            io.BytesIO(classroom_excel_service.build_export_xlsx(result)),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=classroom_schedule_parse.xlsx"})
    return result


@router.delete("/schedule", summary="清空教室课表（不影响个人课表）")
def schedule_clear(db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    from app.services import dashboard_service

    n = _clear_classroom_courses(db, user.id)
    dashboard_service.invalidate(user.id)
    return {"deleted": n}


# ─────────── 使用记录（Excel 导入 · 防重复录入）───────────
@router.post("/usage/import-excel", summary="导入教室使用信息 Excel（同教室同日时间段重叠自动拦截）")
async def usage_import_excel(file: UploadFile = File(...),
                             db: Session = Depends(get_db),
                             user: User = Depends(get_current_user)) -> dict:
    name = (file.filename or "").lower()
    if not name.endswith((".xlsx", ".xls")):
        raise AppError("请选择 .xlsx / .xls 文件")
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise AppError("文件超过 10MB 限制")
    result = classroom_service.import_usage_records(db, data, user_id=user.id)
    from app.services import dashboard_service
    dashboard_service.invalidate(user.id)
    return result


@router.get("/usage-at", summary="某教室在指定时段的使用详情（当前时段 + 上次使用）")
def usage_at(building: str = Query(...), room_no: str = Query(...),
             day: date | None = Query(default=None), hour: int | None = Query(default=None, ge=0, le=23),
             db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    return classroom_service.room_usage_at(db, building, room_no, day=day, hour=hour)


@router.get("/usage/template", summary="下载教室使用信息 Excel 模板")
def usage_template(user: User = Depends(get_current_user)):
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "教室使用信息"
    headers = ["教室编号", "教学楼", "日期", "时间段", "使用状态", "上次使用部门", "用途", "备注"]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1F6FEB")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
    ws.append(["10#101", "10教", "2026-10-05", "08:00-09:40", "使用", "教务处", "高等数学补课", ""])
    ws.append(["10#102", "10教", "2026-10-05", "10:00-11:40", "空闲", "", "", "已消毒"])
    ws.append(["信息#201", "信息楼", "2026-10-06", "14:00-15:40", "使用", "计算机学院", "程序设计实训", ""])
    for col, w in zip("ABCDEFGH", [14, 12, 13, 15, 10, 16, 18, 16]):
        ws.column_dimensions[col].width = w
    ws["A10"] = "填写说明：① 教室编号形如 10#101；② 日期支持 2026-10-05 / 2026/10/5；"
    ws["A11"] = "③ 时间段形如 08:00-09:40（也可拆成开始时间/结束时间两列）；④ 使用状态：使用/空闲；"
    ws["A12"] = "⑤ 同一教室同一天时间段重叠的记录会被判定为重复录入并自动跳过。"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=classroom_usage_template.xlsx"})


# ─────────── 教室信息管理 ───────────
@router.get("/classrooms", summary="教室列表")
def list_classrooms(building: str | None = Query(default=None),
                    room_type: str | None = Query(default=None),
                    keyword: str | None = Query(default=None),
                    only_active: bool = Query(default=True),
                    db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    rooms = classroom_service.list_classrooms(db, building=building, room_type=room_type,
                                              keyword=keyword, only_active=only_active)
    return {"total": len(rooms), "items": [r.to_dict() for r in rooms]}


@router.post("/classrooms", summary="新建教室信息")
def create_classroom(body: ClassroomIn, db: Session = Depends(get_db),
                     user: User = Depends(get_current_user)) -> dict:
    room = classroom_service.get_or_create_classroom(
        db, body.building, body.room_no, capacity=body.capacity, seats=body.seats,
        open_time=body.open_time, close_time=body.close_time, floor=body.floor,
        room_type=body.room_type, has_projector=body.has_projector, has_ac=body.has_ac,
        note=body.note,
    )
    db.commit()
    db.refresh(room)
    return room.to_dict()


@router.get("/classrooms/{classroom_id}", summary="教室详情")
def get_classroom(classroom_id: int, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)) -> dict:
    return classroom_service.get_classroom(db, classroom_id).to_dict()


@router.put("/classrooms/{classroom_id}", summary="更新教室信息")
def update_classroom(classroom_id: int, body: ClassroomUpdate, db: Session = Depends(get_db),
                     user: User = Depends(get_current_user)) -> dict:
    room = classroom_service.get_classroom(db, classroom_id)
    for key, value in body.model_dump(exclude_unset=True).items():
        if value is not None:
            setattr(room, key, value)
    db.commit()
    db.refresh(room)
    return room.to_dict()


@router.delete("/classrooms/{classroom_id}", summary="删除教室（软删除：置为停用）")
def delete_classroom(classroom_id: int, hard: bool = Query(default=False),
                     db: Session = Depends(get_db), user: User = Depends(get_current_user)) -> dict:
    room = classroom_service.get_classroom(db, classroom_id)
    if hard:
        db.delete(room)
    else:
        room.is_active = False
    db.commit()
    return {"deleted": True, "hard": hard}
