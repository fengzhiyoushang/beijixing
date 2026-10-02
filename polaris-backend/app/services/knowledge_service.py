"""知识库服务：文件夹、文档、解析、上传、量化统计。"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.exceptions import AppError, NotFoundError
from app.models.knowledge import KnowledgeChunk, KnowledgeDoc, KnowledgeFolder


# ─────────── 文件夹 ───────────
def list_folders(db: Session, user_id: int) -> list[KnowledgeFolder]:
    return (db.query(KnowledgeFolder)
            .filter(KnowledgeFolder.user_id == user_id)
            .order_by(KnowledgeFolder.sort_order, KnowledgeFolder.id).all())


def create_folder(db: Session, user_id: int, payload) -> KnowledgeFolder:
    if payload.parent_id:
        parent = db.get(KnowledgeFolder, payload.parent_id)
        if not parent or parent.user_id != user_id:
            raise NotFoundError("父文件夹不存在")
    folder = KnowledgeFolder(user_id=user_id, name=payload.name, parent_id=payload.parent_id,
                             icon=payload.icon, sort_order=payload.sort_order)
    db.add(folder)
    db.commit()
    db.refresh(folder)
    return folder


def get_folder(db: Session, user_id: int, folder_id: int) -> KnowledgeFolder:
    folder = db.get(KnowledgeFolder, folder_id)
    if not folder or folder.user_id != user_id:
        raise NotFoundError("文件夹不存在")
    return folder


def update_folder(db: Session, user_id: int, folder_id: int, payload) -> KnowledgeFolder:
    folder = get_folder(db, user_id, folder_id)
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(folder, key, value)
    db.commit()
    db.refresh(folder)
    return folder


def delete_folder(db: Session, user_id: int, folder_id: int, *, delete_docs: bool = False) -> dict:
    folder = get_folder(db, user_id, folder_id)
    doc_count = db.query(KnowledgeDoc).filter(KnowledgeDoc.folder_id == folder.id).count()
    if delete_docs:
        docs = db.query(KnowledgeDoc).filter(KnowledgeDoc.folder_id == folder.id).all()
        for doc in docs:
            db.delete(doc)
    else:
        db.query(KnowledgeDoc).filter(KnowledgeDoc.folder_id == folder.id).update({"folder_id": None})
    db.delete(folder)
    db.commit()
    return {"deleted": True, "docs_affected": doc_count, "docs_deleted": delete_docs}


# ─────────── 文档 ───────────
def create_doc(db: Session, user_id: int, payload) -> KnowledgeDoc:
    doc = KnowledgeDoc(
        user_id=user_id,
        folder_id=payload.folder_id,
        title=payload.title.strip(),
        content=payload.content or "",
        summary=payload.summary or _make_summary(payload.content or ""),
        tags=payload.tags or [],
        doc_type=payload.doc_type,
        word_count=len(payload.content or ""),
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


def get_doc(db: Session, user_id: int, doc_id: int) -> KnowledgeDoc:
    doc = db.get(KnowledgeDoc, doc_id)
    if not doc or doc.user_id != user_id:
        raise NotFoundError("文档不存在")
    return doc


def list_docs(db: Session, user_id: int, *, folder_id: int | None = None,
              keyword: str | None = None, tag: str | None = None,
              vector_status: str | None = None, limit: int = 100) -> list[KnowledgeDoc]:
    query = db.query(KnowledgeDoc).filter(KnowledgeDoc.user_id == user_id)
    if folder_id:
        query = query.filter(KnowledgeDoc.folder_id == folder_id)
    if vector_status:
        query = query.filter(KnowledgeDoc.vector_status == vector_status)
    if keyword:
        like = f"%{keyword}%"
        query = query.filter(KnowledgeDoc.title.like(like) | KnowledgeDoc.content.like(like))
    if tag:
        query = query.filter(KnowledgeDoc.tags.contains([tag]))
    return query.order_by(KnowledgeDoc.updated_at.desc()).limit(limit).all()


def update_doc(db: Session, user_id: int, doc_id: int, payload) -> KnowledgeDoc:
    doc = get_doc(db, user_id, doc_id)
    data = payload.model_dump(exclude_unset=True)
    data.pop("re_vectorize", None)
    for key, value in data.items():
        setattr(doc, key, value)
    if "content" in data and data["content"] is not None:
        doc.word_count = len(doc.content or "")
        doc.summary = doc.summary or _make_summary(doc.content)
    db.commit()
    db.refresh(doc)
    return doc


def delete_doc(db: Session, user_id: int, doc_id: int) -> None:
    doc = get_doc(db, user_id, doc_id)
    db.delete(doc)
    db.commit()


# ─────────── 解析上传文件 ───────────
def extract_text(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext in {"md", "markdown", "txt"}:
        for encoding in ("utf-8", "gbk", "utf-16"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", "ignore")
    if ext == "docx":
        try:
            import docx
        except ImportError:                      # pragma: no cover
            raise AppError("服务端未安装 python-docx，无法解析 .docx")
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if ext == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:                      # pragma: no cover
            raise AppError("服务端未安装 openpyxl，无法解析 .xlsx")
        wb = load_workbook(io.BytesIO(data), data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if all(c is None for c in row):
                    continue
                lines.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(lines)
    raise AppError(f"暂不支持解析 .{ext} 文件，请上传 md/txt/docx/xlsx")


def _make_summary(content: str) -> str:
    text = re.sub(r"[#>*`\-\|]", " ", content or "")
    text = re.sub(r"\s+", " ", text).strip()
    return text[:160]


# ─────────── 量化统计 ───────────
def stats(db: Session, user_id: int) -> dict:
    docs = db.query(KnowledgeDoc).filter(KnowledgeDoc.user_id == user_id).all()
    total_words = sum(d.word_count or 0 for d in docs)
    done = sum(1 for d in docs if d.vector_status == "done")
    chunk_count = (db.query(func.count(KnowledgeChunk.id))
                   .join(KnowledgeDoc, KnowledgeChunk.doc_id == KnowledgeDoc.id)
                   .filter(KnowledgeDoc.user_id == user_id).scalar() or 0)

    # 近 30 天入库节奏
    growth = []
    for i in range(29, -1, -1):
        day = date.today() - timedelta(days=i)
        count = sum(1 for d in docs if d.created_at and d.created_at.date() == day)
        growth.append({"date": day.isoformat(), "count": count})

    tags: dict[str, int] = {}
    for doc in docs:
        for tag in doc.tags or []:
            tags[tag] = tags.get(tag, 0) + 1

    folders = {f.id: f.name for f in list_folders(db, user_id)}
    by_folder: dict[str, int] = {}
    for doc in docs:
        name = folders.get(doc.folder_id, "未分类")
        by_folder[name] = by_folder.get(name, 0) + 1

    top_referenced = sorted(docs, key=lambda d: -(d.read_count or 0))[:5]

    return {
        "doc_count": len(docs),
        "total_words": total_words,
        "chunk_count": int(chunk_count),
        "vector_done": done,
        "vector_pending": sum(1 for d in docs if d.vector_status in {"pending", "processing"}),
        "vector_failed": sum(1 for d in docs if d.vector_status == "failed"),
        "coverage": round(done / len(docs), 4) if docs else 0,
        "new_this_week": sum(1 for d in docs if d.created_at and
                             d.created_at >= datetime.now() - timedelta(days=7)),
        "qa_references_total": sum(d.read_count or 0 for d in docs),
        "avg_words": round(total_words / len(docs), 1) if docs else 0,
        "growth30": growth,
        "hot_tags": [{"tag": k, "count": v} for k, v in sorted(tags.items(), key=lambda x: -x[1])[:8]],
        "by_folder": [{"folder": k, "count": v} for k, v in by_folder.items()],
        "top_referenced": [{"id": d.id, "title": d.title, "read_count": d.read_count or 0}
                           for d in top_referenced],
    }
