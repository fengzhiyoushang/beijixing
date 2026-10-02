"""课程表：CRUD / 冲突检测 / 周课表 / 今日课程 / JSON·Excel 导入。"""
import io
import re
from datetime import datetime

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy.orm import Session

from app import models as m
from app.core.database import get_db
from app.core.deps import get_current_user
from app.schemas import CourseImportIn, CourseIn, CoursePatchIn, SlotIn
from app.services import conflict as conflict_svc
from app.services import dashboard as dash_svc

router = APIRouter(prefix="/courses", tags=["课程表"])

WEEKDAY_MAP = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "日": 7, "天": 7}


def _build_course(db: Session, uid: int, payload: CourseIn) -> m.Course:
    course = m.Course(user_id=uid, name=payload.name, teacher=payload.teacher,
                      location=payload.location, color=payload.color, remark=payload.remark)
    for s in payload.slots:
        course.slots.append(m.ClassSlot(
            weekday=s.weekday, start_time=s.start_time, end_time=s.end_time,
            start_week=s.start_week, end_week=s.end_week, weeks_text=s.weeks_text,
        ))
    db.add(course)
    db.flush()
    return course


@router.get("")
def list_courses(db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    courses = db.query(m.Course).filter(m.Course.user_id == user.id).order_by(m.Course.id).all()
    return [c.to_dict() for c in courses]


@router.post("", status_code=201)
def create_course(
    body: CourseIn,
    force: bool = Query(False, description="忽略冲突强制保存"),
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    if not force:
        conflicts = conflict_svc.find_conflicts(db, user.id, body.slots, course_name=body.name)
        if conflicts:
            raise HTTPException(409, {"message": "检测到课程时间冲突", "conflicts": conflicts})
    course = _build_course(db, user.id, body)
    db.commit()
    db.refresh(course)
    return course.to_dict()


@router.put("/{course_id}")
def update_course(
    course_id: int,
    body: CoursePatchIn,
    force: bool = Query(False),
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    course = db.get(m.Course, course_id)
    if not course or course.user_id != user.id:
        raise HTTPException(404, "课程不存在")
    if body.slots is not None and not force:
        conflicts = conflict_svc.find_conflicts(db, user.id, body.slots,
                                                exclude_course_id=course_id, course_name=course.name)
        if conflicts:
            raise HTTPException(409, {"message": "检测到课程时间冲突", "conflicts": conflicts})
    for field in ("name", "teacher", "location", "color", "remark"):
        value = getattr(body, field)
        if value is not None:
            setattr(course, field, value)
    if body.slots is not None:
        course.slots.clear()
        db.flush()
        for s in body.slots:
            course.slots.append(m.ClassSlot(
                weekday=s.weekday, start_time=s.start_time, end_time=s.end_time,
                start_week=s.start_week, end_week=s.end_week, weeks_text=s.weeks_text,
            ))
    db.commit()
    db.refresh(course)
    return course.to_dict()


@router.delete("/{course_id}")
def delete_course(course_id: int, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    course = db.get(m.Course, course_id)
    if not course or course.user_id != user.id:
        raise HTTPException(404, "课程不存在")
    db.delete(course)
    db.commit()
    return {"deleted": course_id}


@router.post("/check-conflict")
def check_conflict(
    slots: list[SlotIn],
    exclude_course_id: int | None = Query(None),
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    return {"conflicts": conflict_svc.find_conflicts(db, user.id, slots, exclude_course_id)}


@router.get("/week")
def week_grid(db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    days = [{"weekday": wd, "classes": []} for wd in range(1, 8)]
    for c in db.query(m.Course).filter(m.Course.user_id == user.id).all():
        for s in c.slots:
            days[s.weekday - 1]["classes"].append({
                **s.to_dict(), "course_id": c.id, "course": c.name,
                "teacher": c.teacher, "location": c.location, "color": c.color,
            })
    for d in days:
        d["classes"].sort(key=lambda x: x["start_time"])
    return {"days": days}


@router.get("/today")
def today(db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    classes = dash_svc.today_classes(db, user.id)
    return {"date": datetime.now().date().isoformat(), "classes": classes}


@router.post("/import")
def import_courses(
    body: CourseImportIn,
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    """JSON 批量导入：全部写入，同时回报冲突清单。"""
    report = {"imported": 0, "conflicts": []}
    for payload in body.courses:
        found = conflict_svc.find_conflicts(db, user.id, payload.slots, course_name=payload.name)
        for f in found:
            f["importing_course"] = payload.name
        report["conflicts"].extend(found)
        _build_course(db, user.id, payload)
        report["imported"] += 1
    db.commit()
    return report


def _norm_time(value) -> str:
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if hasattr(value, "strftime"):
        return value.strftime("%H:%M")
    text = str(value or "").strip()
    m2 = re.match(r"^(\d{1,2})[:：](\d{2})", text)
    if m2:
        return f"{int(m2.group(1)):02d}:{m2.group(2)}"
    return "00:00"


def _norm_weekday(value) -> int:
    text = str(value or "").strip()
    if text.isdigit():
        wd = int(text)
        return wd if 1 <= wd <= 7 else 1
    for ch, wd in WEEKDAY_MAP.items():
        if ch in text:
            return wd
    return 1


@router.post("/import-excel")
async def import_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    """Excel 导入：列顺序 = 课程名 / 教师 / 地点 / 星期 / 开始 / 结束 / 周次(可选)。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(500, "服务端未安装 openpyxl")
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "无法解析该 Excel 文件")
    ws = wb.active
    courses: dict[str, CourseIn] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if not row or i == 0 and isinstance(row[0], str) and "课程" in (row[0] or ""):
            continue  # 跳过表头
        cells = list(row) + [None] * (7 - len(row))
        name, teacher, location, weekday, start, end, weeks = cells[:7]
        if not name:
            continue
        slot = SlotIn(
            weekday=_norm_weekday(weekday), start_time=_norm_time(start), end_time=_norm_time(end),
            weeks_text=str(weeks) if weeks else None,
        )
        if str(name) in courses:
            courses[str(name)].slots.append(slot)
        else:
            courses[str(name)] = CourseIn(
                name=str(name), teacher=str(teacher) if teacher else None,
                location=str(location) if location else None, slots=[slot],
            )
    report = {"imported": 0, "conflicts": []}
    for payload in courses.values():
        found = conflict_svc.find_conflicts(db, user.id, payload.slots, course_name=payload.name)
        for f in found:
            f["importing_course"] = payload.name
        report["conflicts"].extend(found)
        _build_course(db, user.id, payload)
        report["imported"] += 1
    db.commit()
    return report
