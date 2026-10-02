"""知识库：文档 CRUD / 多格式上传解析 / RAG 问答 / 量化评估。"""
import io
import os
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy import func
from sqlalchemy.orm import Session

from app import models as m
from app.core.database import get_db
from app.core.deps import get_current_user
from app.services import deepseek, rag
from app.services.storage import save_upload
from app.schemas import DocIn, DocPatchIn, QaIn

router = APIRouter(prefix="/knowledge", tags=["知识库"])

SUPPORTED = (".md", ".txt", ".docx", ".xlsx")


def _extract_text(ext: str, content: bytes) -> str:
    if ext in (".md", ".txt"):
        return content.decode("utf-8", errors="ignore")
    if ext == ".docx":
        try:
            import docx
            d = docx.Document(io.BytesIO(content))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        except Exception as exc:
            raise HTTPException(400, f"docx 解析失败: {exc}")
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"# Sheet: {ws.title}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 500:
                        break
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)
        except Exception as exc:
            raise HTTPException(400, f"xlsx 解析失败: {exc}")
    raise HTTPException(400, f"不支持的格式: {ext}")


@router.get("/docs")
def list_docs(search: str | None = Query(None), category: str | None = Query(None),
              db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    q = db.query(m.KnowledgeDoc).filter(m.KnowledgeDoc.user_id == user.id)
    if category:
        q = q.filter(m.KnowledgeDoc.category == category)
    if search:
        q = q.filter(m.KnowledgeDoc.title.contains(search) |
                     m.KnowledgeDoc.content.contains(search))
    docs = q.order_by(m.KnowledgeDoc.updated_at.desc()).all()
    cats = (db.query(m.KnowledgeDoc.category, func.count(m.KnowledgeDoc.id))
            .filter(m.KnowledgeDoc.user_id == user.id)
            .group_by(m.KnowledgeDoc.category).all())
    return {"total": len(docs),
            "docs": [d.brief() for d in docs],
            "categories": [{"category": c or "未分类", "count": n} for c, n in cats]}


@router.post("/docs", status_code=201)
def create_doc(body: DocIn, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    doc = m.KnowledgeDoc(user_id=user.id, title=body.title, content=body.content,
                         category=body.category, tags=body.tags, source="manual",
                         file_type="md")
    db.add(doc)
    db.commit()
    db.refresh(doc)
    chunks = rag.build_index(db, doc)
    return {**doc.brief(), "chunks": chunks}


@router.get("/docs/{doc_id}")
def get_doc(doc_id: int, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    doc = db.get(m.KnowledgeDoc, doc_id)
    if not doc or doc.user_id != user.id:
        raise HTTPException(404, "文档不存在")
    return {**doc.brief(), "content": doc.content}


@router.put("/docs/{doc_id}")
def update_doc(doc_id: int, body: DocPatchIn, db: Session = Depends(get_db),
               user: m.User = Depends(get_current_user)):
    doc = db.get(m.KnowledgeDoc, doc_id)
    if not doc or doc.user_id != user.id:
        raise HTTPException(404, "文档不存在")
    data = body.model_dump(exclude_unset=True)
    for k, v in data.items():
        setattr(doc, k, v)
    db.commit()
    chunks = None
    if "content" in data:
        chunks = rag.build_index(db, doc)
        db.refresh(doc)
    return {**doc.brief(), "chunks": chunks}


@router.delete("/docs/{doc_id}")
def delete_doc(doc_id: int, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    doc = db.get(m.KnowledgeDoc, doc_id)
    if not doc or doc.user_id != user.id:
        raise HTTPException(404, "文档不存在")
    db.delete(doc)
    db.commit()
    return {"deleted": doc_id}


@router.post("/upload", status_code=201)
async def upload_doc(file: UploadFile = File(...), category: str | None = Query(None),
                     db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    content = await file.read()
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in SUPPORTED:
        raise HTTPException(400, f"仅支持 {SUPPORTED}，收到 {ext}")
    text = _extract_text(ext, content)
    if not text.strip():
        raise HTTPException(400, "文件内容为空或无法解析出文本")
    file_url = save_upload(content, file.filename, "knowledge")
    doc = m.KnowledgeDoc(
        user_id=user.id,
        title=os.path.splitext(file.filename or "未命名")[0][:200],
        content=text, category=category, source="upload",
        file_type=ext.lstrip("."), file_url=file_url,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    chunks = rag.build_index(db, doc)
    return {**doc.brief(), "chunks": chunks}


@router.post("/qa")
async def qa(body: QaIn, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    hits = rag.search(db, user.id, body.question, top_k=body.top_k)
    for h in hits:
        doc = db.get(m.KnowledgeDoc, h["doc_id"])
        if doc:
            doc.read_count = (doc.read_count or 0) + 1
            doc.last_read_at = datetime.now()
    db.commit()

    references = [{"doc_id": h["doc_id"], "doc_title": h["doc_title"],
                   "snippet": h["text"][:120], "score": h["score"]} for h in hits]
    if not hits:
        return {"question": body.question,
                "answer": "知识库中没有检索到相关内容。可在「知识库」页上传笔记或文档后再试。",
                "references": [], "mode": "empty"}

    if deepseek.is_configured():
        try:
            msg = await deepseek.chat(
                [{"role": "system", "content": "你是知识库问答助手，中文简洁回答。"},
                 {"role": "user", "content": rag.build_qa_prompt(body.question, hits)}],
                temperature=0.3)
            return {"question": body.question, "answer": msg.get("content") or "（无内容）",
                    "references": references, "mode": "deepseek"}
        except Exception as exc:
            fallback = "\n\n".join(f"【{h['doc_title']}】{h['text'][:200]}" for h in hits[:3])
            return {"question": body.question,
                    "answer": f"（AI 调用失败，返回本地检索结果：{exc}）\n\n{fallback}",
                    "references": references, "mode": "extractive"}

    answer = "（降级模式：未配置 DeepSeek Key，以下为检索到的原文片段）\n\n" + \
        "\n\n".join(f"【资料{i+1}·{h['doc_title']}】{h['text'][:200]}" for i, h in enumerate(hits[:3]))
    return {"question": body.question, "answer": answer,
            "references": references, "mode": "extractive"}


@router.get("/stats")
def stats(db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    docs = db.query(m.KnowledgeDoc).filter(m.KnowledgeDoc.user_id == user.id).all()
    total_words = sum(d.word_count or 0 for d in docs)
    chunk_count = (db.query(func.count(m.KnowledgeChunk.id))
                   .join(m.KnowledgeDoc, m.KnowledgeChunk.doc_id == m.KnowledgeDoc.id)
                   .filter(m.KnowledgeDoc.user_id == user.id).scalar()) or 0
    doc_ids = [d.id for d in docs]
    chunked = (db.query(func.count(func.distinct(m.KnowledgeChunk.doc_id)))
               .filter(m.KnowledgeChunk.doc_id.in_(doc_ids)).scalar()) or 0 if doc_ids else 0
    qa_refs = sum(d.read_count or 0 for d in docs)
    today = date.today()
    growth = []
    for i in range(29, -1, -1):
        d = today - timedelta(days=i)
        growth.append({"date": d.isoformat(),
                       "count": sum(1 for doc in docs
                                    if doc.created_at and doc.created_at.date() == d)})
    return {
        "doc_count": len(docs),
        "total_words": total_words,
        "avg_words": round(total_words / len(docs)) if docs else 0,
        "chunk_count": chunk_count,
        "coverage": round(chunked / len(docs), 3) if docs else 0,
        "qa_references_total": qa_refs,
        "top_referenced": sorted([{"title": d.title, "read_count": d.read_count or 0} for d in docs],
                                  key=lambda x: -x["read_count"])[:5],
        "growth30": growth,
        "by_category": {},
    }
